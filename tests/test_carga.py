import logging
import sqlite3
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook
import pandas as pd

from src.carga import Carga, ResumenMensualCalendario


class CargaTests(unittest.TestCase):
    def test_carga_y_verifica_sqlite_y_exporta_xlsx(self):
        with tempfile.TemporaryDirectory() as temporary_dir:
            root = Path(temporary_dir)
            loader = Carga(
                sqlite_path=root / "airbnb.sqlite",
                excel_dir=root / "excel",
                logger=logging.getLogger("test"),
            )
            frames = {
                "Listings": pd.DataFrame({"id": [1, 2], "price": [100, 250]}),
                "Reviews": pd.DataFrame({"id": [10], "date": ["2025-01-01"]}),
            }

            loaded = loader.cargar(frames)

            self.assertEqual(loaded, {"listings": 2, "reviews": 1})
            connection = sqlite3.connect(root / "airbnb.sqlite")
            try:
                self.assertEqual(
                    connection.execute("SELECT COUNT(*) FROM listings").fetchone()[0],
                    2,
                )
            finally:
                connection.close()
            self.assertEqual(
                pd.read_excel(root / "excel" / "listings.xlsx")["price"].tolist(),
                [100, 250],
            )

    def test_rechaza_dataframe_sin_columnas(self):
        with tempfile.TemporaryDirectory() as temporary_dir:
            loader = Carga(
                sqlite_path=Path(temporary_dir) / "airbnb.sqlite",
                excel_dir=Path(temporary_dir) / "excel",
                logger=logging.getLogger("test"),
            )

            with self.assertRaisesRegex(ValueError, "sin columnas"):
                loader.cargar({"Listings": pd.DataFrame()})

    def test_carga_bloques_y_divide_xlsx_en_hojas(self):
        with tempfile.TemporaryDirectory() as temporary_dir:
            root = Path(temporary_dir)
            loader = Carga(
                sqlite_path=root / "airbnb.sqlite",
                excel_dir=root / "excel",
                logger=logging.getLogger("test"),
            )
            loader.EXCEL_MAX_ROWS = 3
            blocks = [
                pd.DataFrame({"id": [1, 2], "price": [100, 200]}),
                pd.DataFrame({"id": [3, 4, 5], "price": [300, 400, 500]}),
            ]

            loaded = loader.cargar_bloques("Calendar", blocks)

            self.assertEqual(loaded, 5)
            workbook = load_workbook(
                root / "excel" / "calendar.xlsx", read_only=True
            )
            try:
                self.assertEqual(workbook.sheetnames, ["calendar_001", "calendar_002", "calendar_003"])
                self.assertEqual(
                    list(workbook["calendar_003"].values),
                    [("id", "price"), (5, 500)],
                )
            finally:
                workbook.close()

    def test_exporta_resumen_de_calendar_como_calendar_xlsx(self):
        with tempfile.TemporaryDirectory() as temporary_dir:
            root = Path(temporary_dir)
            loader = Carga(
                sqlite_path=root / "airbnb.sqlite",
                excel_dir=root / "excel",
                logger=logging.getLogger("test"),
            )
            resumen = ResumenMensualCalendario()
            for _ in resumen.observar(
                [
                    pd.DataFrame(
                        {
                            "date_year": [2025, 2025, 2025],
                            "date_month": [1, 1, 2],
                            "available": [True, False, True],
                            "listing_id": [10, 10, 20],
                        }
                    )
                ]
            ):
                pass

            path = loader.exportar_resumen_calendario(resumen.dataframe())

            self.assertEqual(path.name, "calendar.xlsx")
            result = pd.read_excel(path)
            self.assertEqual(result["anio"].tolist(), [2025, 2025])
            self.assertEqual(result["mes"].tolist(), [1, 2])
            self.assertEqual(result["dias_disponibles"].tolist(), [1, 1])
            self.assertEqual(result["anuncios_distintos"].tolist(), [1, 1])


if __name__ == "__main__":
    unittest.main()
